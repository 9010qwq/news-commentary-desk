from __future__ import annotations
import hashlib,json,re,uuid,zipfile
from datetime import date,timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from .store import now

def days(start,end):
    a,b=date.fromisoformat(start),date.fromisoformat(end)
    if a>b or (b-a).days>30:raise ValueError('范围需为1至31天，开始日期不能晚于结束日期')
    return [(a+timedelta(days=i)).isoformat() for i in range((b-a).days+1)]
def safe_cell(value):
    s=str(value or '')
    s=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',s)
    return "'"+s if s.startswith(('=','+','-','@')) else s
def safe_filename(value):
    s=re.sub(r'[\\/:*?"<>|\x00-\x1f]','_',str(value)).strip(' .')[:60] or 'media'
    if s.upper() in {'CON','PRN','AUX','NUL',*[f'COM{i}' for i in range(1,10)],*[f'LPT{i}' for i in range(1,10)]}:s='_'+s
    return s

def export_package(store,start,end):
    period=days(start,end);settings=store.settings();articles=store.articles(start,end)
    suffix=uuid.uuid4().hex[:8];stem=f'新闻时评_{start}_{end}_{suffix}'
    xlsx=store.root/'exports'/(stem+'.xlsx');zpath=store.root/'exports'/(stem+'_截图.zip')
    wb=Workbook();ws=wb.active;ws.title='新闻时评'
    headers=['日期（刊发）','标题','作者（原页署名）','媒体','原文链接','原页发布时间','实际截图时间','截图状态','截图文件（ZIP内）','时评判定','原页来源','问题说明']
    ws.append(headers)
    manifest={'range':{'start':start,'end':end},'created_at':now(),'timezone':settings.timezone,'date_basis':'原文刊发日期；截图时间为实际采集时刻','items':[],'daily':[]}
    with zipfile.ZipFile(zpath,'w',zipfile.ZIP_DEFLATED) as z:
        for a in articles:
            status=a['screenshot_status'];relative='';error=a['error']
            if status=='ok':
                p=(store.root/a['screenshot_path']).resolve()
                if p.is_relative_to((store.root/'screenshots').resolve()) and p.is_file():
                    relative=f"{a['date']}/{safe_filename(a['media'])}_{a['id'][:10]}.png";z.write(p,relative)
                else:status='missing';error='截图文件缺失，未加入ZIP'
            row=[a['date'],a['title'],a['author'],a['media'],a['url'],a['published_at'],a['captured_at'],status,relative,a['classification'],a['original_source'],error]
            ws.append([safe_cell(v) for v in row]);ws.cell(ws.max_row,5).hyperlink=a['url'];ws.cell(ws.max_row,5).style='Hyperlink'
            manifest['items'].append({k:v for k,v in a.items() if k not in {'title_hash'}}|{'screenshot_status':status,'zip_path':relative,'sha256':hashlib.sha256(p.read_bytes()).hexdigest() if relative else None,'error':error})
        for d in period:
            rows=[a for a in manifest['items'] if a['date']==d];complete=sum(a['screenshot_status']=='ok' for a in rows)
            manifest['daily'].append({'date':d,'target':settings.target_count,'articles':len(rows),'screenshots':complete,'shortfall':max(0,settings.target_count-complete)})
        z.writestr('manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2))
        z.writestr('说明.txt','这是原媒体网页的实际浏览器截图；不是文字卡片。按刊发日期分文件夹。补采截图时间不等于历史刊发日期。缺失、拒访、截图失败均记录在 manifest.json 和 Excel，未编造补足。请遵守原媒体版权及使用限制。')
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    widths=[15,65,26,18,62,28,28,16,55,25,22,50]
    for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='23523E');c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=30
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
        ws.row_dimensions[row[0].row].height=56
    summary=wb.create_sheet('每日缺口');summary.append(['日期','每日目标','真实条目','成功截图','尚缺截图','说明'])
    for d in manifest['daily']:summary.append([d['date'],d['target'],d['articles'],d['screenshots'],d['shortfall'],'未达标时不编造；可补充原文链接再采集' if d['shortfall'] else '已达到目标'])
    for i,w in enumerate([17,15,15,15,15,55],1):summary.column_dimensions[get_column_letter(i)].width=w
    for c in summary[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='23523E')
    summary.freeze_panes='A2';summary.auto_filter.ref=summary.dimensions
    summary.append([]);summary.append(['日期范围',start+' 至 '+end]);summary.append(['生成时间',manifest['created_at']]);summary.append(['时区',settings.timezone]);summary.append(['口径','刊发日期归档；采集/截图时间单独记录。缺失不计成功，原页未署名记“未署名”。'])
    wb.save(xlsx)
    return {'files':[{'name':p.name,'url':'/files/exports/'+p.name,'path':str(p)} for p in (xlsx,zpath)],'counts':{'articles':len(articles),'screenshots':sum(x['screenshots'] for x in manifest['daily'])},'shortfalls':manifest['daily']}
